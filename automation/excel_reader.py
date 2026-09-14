from pathlib import Path
from openpyxl import load_workbook

from models import TestCase
from config import EXCEL_FILE

REQUIRED_COLUMNS = [
    "Test ID",
    "Test Name",
    "Category",
    "Steps",
    "Expected Result",
    "Status",
    "AI Observation",
]


def load_test_cases() -> list[TestCase]:
    excel_path = Path(EXCEL_FILE)
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found at {excel_path}")

    wb = load_workbook(excel_path)
    sheet = wb.active

    headers = [cell.value for cell in sheet[1]]
    missing = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise ValueError(
            f"Missing required column(s): {', '.join(missing)}. "
            f"Found columns: {headers}"
        )

    col_idx = {c: headers.index(c) for c in REQUIRED_COLUMNS}
    test_cases: list[TestCase] = []

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        test_id = row[col_idx["Test ID"]]
        if test_id is None or str(test_id).strip() == "":
            continue  # skip blank rows

        raw_steps = row[col_idx["Steps"]]
        if raw_steps is None or str(raw_steps).strip() == "":
            raise ValueError(
                f"Row {row_number}: Test '{test_id}' has empty Steps column."
            )

        test_cases.append(
            TestCase(
                test_id=str(test_id).strip(),
                test_name=str(row[col_idx["Test Name"]] or "").strip(),
                category=str(row[col_idx["Category"]] or "").strip(),
                raw_steps=str(raw_steps),
                expected_result=str(row[col_idx["Expected Result"]] or "").strip(),
                status=str(row[col_idx["Status"]] or "NOT RUN").strip(),
                ai_observation=(
                    str(row[col_idx["AI Observation"]]).strip()
                    if row[col_idx["AI Observation"]]
                    else None
                ),
            )
        )

    if not test_cases:
        raise ValueError("No test cases were found in the Excel file.")

    return test_cases